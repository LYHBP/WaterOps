"""防汛物资流水、库存、权限与制式台账导出。"""

from __future__ import annotations

from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

from alembic import command
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import inspect, text

from app.config import Settings
from app.database import DatabaseRuntime


def _login(client: TestClient, username: str) -> None:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": "PartyOps@2026"},
    )
    assert response.status_code == 200, response.text


def _entry(
    client: TestClient,
    item_id: str,
    kind: str,
    quantity: int,
    **extra,
) -> dict:
    payload = {
        "item_id": item_id,
        "transaction_type": kind,
        "quantity": quantity,
        "occurred_at": "2026-09-30T09:30:00+08:00",
        "document_no": f"TEST-{kind}-{quantity}",
    }
    payload.update(extra)
    response = client.post("/api/v1/inventory/transactions", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def test_inventory_lifecycle_stock_forms_and_idempotency(
    client: TestClient, admin: dict
) -> None:
    _login(client, "admin")
    created = client.post(
        "/api/v1/inventory/items",
        json={
            "name": "闭环测试救生衣",
            "specification": "成人款",
            "unit": "件",
            "storage_location": "一号库",
            "category": "flood_control",
        },
    )
    assert created.status_code == 201, created.text
    item = created.json()

    insufficient = client.post(
        "/api/v1/inventory/transactions",
        json={
            "item_id": item["id"],
            "transaction_type": "outbound",
            "quantity": 1,
            "occurred_at": "2026-09-30T09:30:00+08:00",
        },
    )
    assert insufficient.status_code == 409
    assert insufficient.json()["code"] == "INVENTORY_INSUFFICIENT_STOCK"

    inbound_payload = {
        "item_id": item["id"],
        "transaction_type": "inbound",
        "quantity": 10,
        "occurred_at": "2026-09-30T09:30:00+08:00",
        "document_no": "RK-TEST-001",
        "client_request_id": "inventory-inbound-test-001",
    }
    inbound = client.post("/api/v1/inventory/transactions", json=inbound_payload)
    assert inbound.status_code == 201, inbound.text
    repeated = client.post("/api/v1/inventory/transactions", json=inbound_payload)
    assert repeated.status_code == 201, repeated.text
    assert repeated.json()["id"] == inbound.json()["id"]

    reused_key = client.post(
        "/api/v1/inventory/transactions",
        json={**inbound_payload, "quantity": 9},
    )
    assert reused_key.status_code == 409
    assert reused_key.json()["code"] == "IDEMPOTENCY_KEY_REUSED"

    loan = _entry(
        client,
        item["id"],
        "loan",
        4,
        expected_return_at="2026-10-05T17:00:00+08:00",
        counterparty="防汛值守组",
    )
    outbound = _entry(client, item["id"], "outbound", 2, purpose="应急领用")
    returned_first = _entry(
        client,
        item["id"],
        "return",
        1,
        related_transaction_id=loan["id"],
    )
    _entry(
        client,
        item["id"],
        "return",
        2,
        related_transaction_id=loan["id"],
    )
    over_return = client.post(
        "/api/v1/inventory/transactions",
        json={
            "item_id": item["id"],
            "transaction_type": "return",
            "quantity": 2,
            "occurred_at": "2026-09-30T09:30:00+08:00",
            "related_transaction_id": loan["id"],
        },
    )
    assert over_return.status_code == 409
    assert over_return.json()["code"] == "RETURN_EXCEEDS_LOAN"

    missing_adjustment_reason = client.post(
        "/api/v1/inventory/transactions",
        json={
            "item_id": item["id"],
            "transaction_type": "adjustment_out",
            "quantity": 1,
            "occurred_at": "2026-09-30T09:30:00+08:00",
        },
    )
    assert missing_adjustment_reason.status_code == 422
    adjustment = _entry(client, item["id"], "adjustment_out", 1, purpose="盘点差异：实物少1件")
    stock = client.get("/api/v1/inventory/stock")
    assert stock.status_code == 200, stock.text
    row = next(value for value in stock.json() if value["item"]["id"] == item["id"])
    assert row["inbound_quantity"] == 10
    assert row["loaned_quantity"] == 4
    assert row["returned_quantity"] == 3
    assert row["adjustment_quantity"] == -1
    assert row["available_quantity"] == 6

    for entry, form in ((inbound.json(), "M-01"), (outbound, "M-02")):
        exported_form = client.get(f"/api/v1/inventory/transactions/{entry['id']}/{form}.xlsx")
        assert exported_form.status_code == 200, exported_form.text
        workbook = load_workbook(BytesIO(exported_form.content), read_only=True)
        assert workbook.active.max_row >= 10
        assert "待甲方样表确认" in str(workbook.active[1][0].value)

    ledger_form = client.get("/api/v1/inventory/forms/M-05.xlsx")
    assert ledger_form.status_code == 200, ledger_form.text
    ledger = load_workbook(BytesIO(ledger_form.content), read_only=True, data_only=False)
    assert ledger["库存及出入库台账"]["K2"].value == row["available_quantity"]

    exported_stock = client.get("/api/v1/inventory/export.xlsx")
    assert exported_stock.status_code == 200, exported_stock.text
    workbook = load_workbook(BytesIO(exported_stock.content), read_only=True, data_only=False)
    sheet = workbook["库存及出入库台账"]
    data_row = next(row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0] == item["name"])
    assert data_row[10] == row["available_quantity"]

    patched = client.patch(
        f"/api/v1/inventory/items/{item['id']}",
        headers={"If-Match": str(item["version"])},
        json={"storage_location": "二号库"},
    )
    assert patched.status_code == 200, patched.text
    stale = client.patch(
        f"/api/v1/inventory/items/{item['id']}",
        headers={"If-Match": str(item["version"])},
        json={"storage_location": "错误库位"},
    )
    assert stale.status_code == 409
    assert stale.json()["code"] == "VERSION_CONFLICT"

    voided = client.post(
        f"/api/v1/inventory/transactions/{adjustment['id']}/void",
        headers={"If-Match": str(adjustment["version"])},
        json={"reason": "测试盘点调整录入错误"},
    )
    assert voided.status_code == 200, voided.text
    assert voided.json()["status"] == "voided"
    after_void = client.get("/api/v1/inventory/stock").json()
    row = next(value for value in after_void if value["item"]["id"] == item["id"])
    assert row["available_quantity"] == 7

    void_inbound = client.post(
        f"/api/v1/inventory/transactions/{inbound.json()['id']}/void",
        headers={"If-Match": str(inbound.json()["version"])},
        json={"reason": "不得形成负库存"},
    )
    assert void_inbound.status_code == 409
    assert void_inbound.json()["code"] == "INVENTORY_INSUFFICIENT_STOCK"


def test_inventory_writes_are_admin_only(
    client: TestClient, admin: dict, staff: dict
) -> None:
    _login(client, "staff")
    response = client.post(
        "/api/v1/inventory/items",
        json={"name": "无权限测试物资", "unit": "件"},
    )
    assert response.status_code == 403
    _login(client, "admin")


def test_concurrent_outbound_cannot_create_negative_stock(
    client: TestClient, admin: dict
) -> None:
    _login(client, "admin")
    created = client.post(
        "/api/v1/inventory/items",
        json={"name": "并发库存测试物资", "unit": "件"},
    )
    assert created.status_code == 201, created.text
    item_id = created.json()["id"]
    _entry(client, item_id, "inbound", 1)
    payload = {
        "item_id": item_id,
        "transaction_type": "outbound",
        "quantity": 1,
        "occurred_at": "2026-09-30T09:30:00+08:00",
    }

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(
            lambda _: client.post("/api/v1/inventory/transactions", json=payload),
            range(2),
        ))

    assert sorted(response.status_code for response in responses) == [201, 409]
    stock = client.get("/api/v1/inventory/stock").json()
    row = next(value for value in stock if value["item"]["id"] == item_id)
    assert row["available_quantity"] == 0


def test_inventory_migration_up_and_down_from_0024(tmp_path: Path) -> None:
    runtime = DatabaseRuntime(Settings(data_dir=tmp_path, environment="test", strict_sqlite=False))
    with runtime.engine.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE users (id VARCHAR(36) PRIMARY KEY)")
        connection.exec_driver_sql(
            "CREATE TABLE alembic_version (version_num VARCHAR(32) NOT NULL PRIMARY KEY)"
        )
        connection.exec_driver_sql("INSERT INTO alembic_version(version_num) VALUES ('0024')")

    runtime._upgrade_to_head()
    with runtime.engine.connect() as connection:
        assert connection.execute(text("SELECT version_num FROM alembic_version")).scalar_one() == "0025"
        assert {"inventory_items", "inventory_transactions"} <= set(inspect(connection).get_table_names())

    with runtime.engine.begin() as connection:
        config = runtime._alembic_config()
        config.attributes["connection"] = connection
        command.downgrade(config, "0024")
    with runtime.engine.connect() as connection:
        assert connection.execute(text("SELECT version_num FROM alembic_version")).scalar_one() == "0024"
        assert not (
            {"inventory_items", "inventory_transactions"}
            & set(inspect(connection).get_table_names())
        )
    runtime.engine.dispose()
